"""excel.py - one function that turns headers + rows into a styled .xlsx download."""
from datetime import date, datetime
from io import BytesIO
from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILLS = {"row-green": "C6EFCE", "row-yellow": "FFEB9C", "row-red": "FFC7CE"}


def xlsx_response(filename, headers, rows, sheet="Data", row_classes=None, money_cols=(), title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    r0 = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        r0 = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows):
        fill = FILLS.get((row_classes or [None] * (i + 1))[i])
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r0 + 1 + i, column=c, value=v)
            if isinstance(v, (date, datetime)):
                cell.number_format = "DD-MM-YYYY"
            elif isinstance(v, float) or (c - 1) in money_cols:
                cell.number_format = "#,##0.00"
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    for c, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[c - 1])) for r in rows[:200] if c - 1 < len(r) and r[c - 1] is not None])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 2), 48)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if rows:
        ws.auto_filter.ref = "A%d:%s%d" % (r0, get_column_letter(len(headers)), r0 + len(rows))
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="%s"' % filename})


def module_export(module, user, args, filename=None):
    """Export a module's list honouring the same filters the user sees on screen."""
    if module.export_override:
        headers, rows = module.export_override(user, args)
        classes = None
    else:
        data = module.query(user, filters=args, sort=args.get("sort"), direction=args.get("dir", "asc"))
        fields = module.export_fields
        headers = [f.label for f in fields]
        rows = [[(row.get(f.name + "__label") if f.kind == "ref" else f.excel(row.get(f.name))) for f in fields]
                for row in data["rows"]]
        classes = [module.row_class(r) if module.row_class else None for r in data["rows"]]
    return xlsx_response(filename or "%s.xlsx" % module.slug, headers, rows, sheet=module.plural,
                         row_classes=classes)
